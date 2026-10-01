import os
import mysql.connector

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference


# ============================================================
# LOGIX — Optimized Executive Dashboard
# Direct MySQL → Dashboard
# ============================================================

OUTPUT_FILE = os.path.join(
    "excel",
    "LOGIX_Executive_Dashboard_v2.xlsx"
)


# ============================================================
# DATABASE CONFIGURATION
# ============================================================

DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "avanish@123",
    "database": "logix"
}


# ============================================================
# THEME
# ============================================================

DARK_BLUE = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"


# ============================================================
# STYLES
# ============================================================

title_fill = PatternFill(
    fill_type="solid",
    fgColor=DARK_BLUE
)

section_fill = PatternFill(
    fill_type="solid",
    fgColor=BLUE
)

kpi_fill = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BLUE
)

white_font = Font(
    color=WHITE,
    bold=True
)

title_font = Font(
    color=WHITE,
    bold=True,
    size=20
)

kpi_font = Font(
    bold=True,
    size=13
)

thin_side = Side(
    style="thin",
    color="D9E1F2"
)

thin_border = Border(
    left=thin_side,
    right=thin_side,
    top=thin_side,
    bottom=thin_side
)


# ============================================================
# MYSQL HELPER
# ============================================================

def fetch_all(cursor, query):
    cursor.execute(query)
    return cursor.fetchall()


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("LOGIX — OPTIMIZED EXECUTIVE DASHBOARD")
    print("=" * 70)

    # ========================================================
    # MYSQL CONNECTION
    # ========================================================

    print("\nConnecting to MySQL...")

    connection = mysql.connector.connect(
        host=DB_CONFIG["host"],
        user=DB_CONFIG["user"],
        password=DB_CONFIG["password"],
        database=DB_CONFIG["database"]
    )

    cursor = connection.cursor()

    print("MySQL connection successful.")

    # ========================================================
    # 1. EXECUTIVE KPI
    # ========================================================

    print("\nFetching Executive KPIs...")

    kpi_query = """
    SELECT
        total_orders,
        total_customers,
        total_revenue,
        total_logistics_cost,
        failed_shipments,
        total_returns,
        total_refunds,
        on_time_rate
    FROM vw_executive_kpi
    """

    cursor.execute(kpi_query)

    kpi = cursor.fetchone()

    if not kpi:
        raise ValueError(
            "No data returned from vw_executive_kpi"
        )

    (
        total_orders,
        total_customers,
        total_revenue,
        total_logistics_cost,
        failed_shipments,
        total_returns,
        total_refunds,
        on_time_rate
    ) = kpi

    # ========================================================
    # 2. MONTHLY SALES
    # ========================================================

    print("Fetching monthly sales...")

    monthly_query = """
    SELECT
        order_month,
        total_orders,
        revenue
    FROM vw_monthly_sales
    ORDER BY order_month
    """

    monthly_data = fetch_all(
        cursor,
        monthly_query
    )

    # ========================================================
    # 3. CUSTOMER SEGMENT REVENUE
    # ========================================================

    print("Fetching customer segment revenue...")

    segment_query = """
    SELECT
        customer_segment,
        ROUND(SUM(line_total), 2) AS revenue
    FROM vw_sales_analysis
    GROUP BY customer_segment
    ORDER BY revenue DESC
    """

    segment_data = fetch_all(
        cursor,
        segment_query
    )

    # ========================================================
    # 4. CATEGORY REVENUE
    # ========================================================

    print("Fetching category revenue...")

    category_query = """
    SELECT
        category,
        ROUND(SUM(line_total), 2) AS revenue
    FROM vw_sales_analysis
    GROUP BY category
    ORDER BY revenue DESC
    """

    category_data = fetch_all(
        cursor,
        category_query
    )

    # ========================================================
    # 5. DELIVERY PERFORMANCE
    # ========================================================

    print("Fetching delivery performance...")

    delivery_query = """
    SELECT
        delivery_performance,
        COUNT(*) AS shipment_count
    FROM vw_delivery_performance
    GROUP BY delivery_performance
    ORDER BY shipment_count DESC
    """

    delivery_data = fetch_all(
        cursor,
        delivery_query
    )

    # ========================================================
    # 6. WAREHOUSE PERFORMANCE
    # ========================================================

    print("Fetching warehouse performance...")

    warehouse_query = """
    SELECT
        warehouse_name,
        total_shipments,
        total_shipping_cost,
        on_time_rate
    FROM vw_warehouse_performance
    ORDER BY total_shipments DESC
    """

    warehouse_data = fetch_all(
        cursor,
        warehouse_query
    )

    # ========================================================
    # CLOSE MYSQL
    # ========================================================

    cursor.close()
    connection.close()

    print("MySQL connection closed.")

    # ========================================================
    # CREATE WORKBOOK
    # ========================================================

    print("\nCreating dashboard workbook...")

    wb = Workbook()

    dashboard = wb.active
    dashboard.title = "Executive_Dashboard"

    data_ws = wb.create_sheet("Chart_Data")

    # ========================================================
    # DASHBOARD SETTINGS
    # ========================================================

    dashboard.sheet_view.showGridLines = False

    for col in range(1, 13):

        dashboard.column_dimensions[
            chr(64 + col)
        ].width = 15

    # ========================================================
    # TITLE
    # ========================================================

    dashboard.merge_cells("A1:L2")

    dashboard["A1"] = (
        "LOGIX — LOGISTICS & SUPPLY CHAIN "
        "INTELLIGENCE DASHBOARD"
    )

    dashboard["A1"].fill = title_fill
    dashboard["A1"].font = title_font

    dashboard["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    dashboard.row_dimensions[1].height = 30
    dashboard.row_dimensions[2].height = 30

    # ========================================================
    # SUBTITLE
    # ========================================================

    dashboard.merge_cells("A3:L3")

    dashboard["A3"] = (
        "Executive Business Performance Overview"
    )

    dashboard["A3"].font = Font(
        italic=True,
        size=10
    )

    dashboard["A3"].alignment = Alignment(
        horizontal="center"
    )

    # ========================================================
    # KPI CARDS
    # ========================================================

    kpis = [
        (
            "A5:C7",
            "Total Orders",
            f"{total_orders:,}"
        ),
        (
            "D5:F7",
            "Total Customers",
            f"{total_customers:,}"
        ),
        (
            "G5:I7",
            "Total Revenue",
            f"₹{total_revenue:,.2f}"
        ),
        (
            "J5:L7",
            "Logistics Cost",
            f"₹{total_logistics_cost:,.2f}"
        ),
        (
            "A9:C11",
            "Failed Shipments",
            f"{failed_shipments:,}"
        ),
        (
            "D9:F11",
            "Total Returns",
            f"{total_returns:,}"
        ),
        (
            "G9:I11",
            "Total Refunds",
            f"₹{total_refunds:,.2f}"
        ),
        (
            "J9:L11",
            "On-Time Delivery",
            f"{on_time_rate:.2f}%"
        )
    ]

    for cell_range, label, value in kpis:

        dashboard.merge_cells(cell_range)

        cell = dashboard[
            cell_range.split(":")[0]
        ]

        cell.value = (
            f"{label}\n"
            f"{value}"
        )

        cell.fill = kpi_fill
        cell.font = kpi_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    # ========================================================
    # SECTION HEADERS
    # ========================================================

    sections = [
        ("A13:F13", "SALES PERFORMANCE"),
        ("G13:L13", "DELIVERY PERFORMANCE"),
        ("A29:F29", "CUSTOMER & REVENUE"),
        ("G29:L29", "OPERATIONS")
    ]

    for cell_range, title in sections:

        dashboard.merge_cells(cell_range)

        cell = dashboard[
            cell_range.split(":")[0]
        ]

        cell.value = title
        cell.fill = section_fill
        cell.font = white_font

        cell.alignment = Alignment(
            horizontal="center"
        )

    # ========================================================
    # CHART DATA — MONTHLY SALES
    # ========================================================

    data_ws["A1"] = "Month"
    data_ws["B1"] = "Orders"
    data_ws["C1"] = "Revenue"

    for row_index, row in enumerate(
        monthly_data,
        start=2
    ):

        data_ws.cell(
            row_index,
            1,
            row[0]
        )

        data_ws.cell(
            row_index,
            2,
            row[1]
        )

        data_ws.cell(
            row_index,
            3,
            float(row[2])
        )

    # ========================================================
    # CHART DATA — CUSTOMER SEGMENT
    # ========================================================

    data_ws["E1"] = "Customer Segment"
    data_ws["F1"] = "Revenue"

    for row_index, row in enumerate(
        segment_data,
        start=2
    ):

        data_ws.cell(
            row_index,
            5,
            row[0]
        )

        data_ws.cell(
            row_index,
            6,
            float(row[1])
        )

    # ========================================================
    # CHART DATA — CATEGORY
    # ========================================================

    data_ws["H1"] = "Category"
    data_ws["I1"] = "Revenue"

    for row_index, row in enumerate(
        category_data,
        start=2
    ):

        data_ws.cell(
            row_index,
            8,
            row[0]
        )

        data_ws.cell(
            row_index,
            9,
            float(row[1])
        )

    # ========================================================
    # CHART DATA — DELIVERY
    # ========================================================

    data_ws["K1"] = "Delivery Performance"
    data_ws["L1"] = "Shipments"

    for row_index, row in enumerate(
        delivery_data,
        start=2
    ):

        data_ws.cell(
            row_index,
            11,
            row[0]
        )

        data_ws.cell(
            row_index,
            12,
            row[1]
        )

    # ========================================================
    # CHART DATA — WAREHOUSE
    # ========================================================

    data_ws["N1"] = "Warehouse"
    data_ws["O1"] = "Shipments"
    data_ws["P1"] = "Shipping Cost"

    for row_index, row in enumerate(
        warehouse_data,
        start=2
    ):

        data_ws.cell(
            row_index,
            14,
            row[0]
        )

        data_ws.cell(
            row_index,
            15,
            row[1]
        )

        data_ws.cell(
            row_index,
            16,
            float(row[2])
        )

    # ========================================================
    # CHART 1 — MONTHLY REVENUE
    # ========================================================

    revenue_chart = LineChart()

    revenue_chart.title = "Monthly Revenue Trend"
    revenue_chart.y_axis.title = "Revenue (₹)"
    revenue_chart.x_axis.title = "Month"

    revenue_data = Reference(
        data_ws,
        min_col=3,
        min_row=1,
        max_row=len(monthly_data) + 1
    )

    revenue_categories = Reference(
        data_ws,
        min_col=1,
        min_row=2,
        max_row=len(monthly_data) + 1
    )

    revenue_chart.add_data(
        revenue_data,
        titles_from_data=True
    )

    revenue_chart.set_categories(
        revenue_categories
    )

    revenue_chart.height = 7
    revenue_chart.width = 12

    dashboard.add_chart(
        revenue_chart,
        "A14"
    )

    # ========================================================
    # CHART 2 — MONTHLY ORDERS
    # ========================================================

    orders_chart = LineChart()

    orders_chart.title = "Monthly Orders Trend"
    orders_chart.y_axis.title = "Orders"
    orders_chart.x_axis.title = "Month"

    orders_data = Reference(
        data_ws,
        min_col=2,
        min_row=1,
        max_row=len(monthly_data) + 1
    )

    orders_chart.add_data(
        orders_data,
        titles_from_data=True
    )

    orders_chart.set_categories(
        revenue_categories
    )

    orders_chart.height = 7
    orders_chart.width = 12

    dashboard.add_chart(
        orders_chart,
        "G14"
    )

    # ========================================================
    # CHART 3 — CUSTOMER SEGMENT REVENUE
    # ========================================================

    segment_chart = BarChart()

    segment_chart.type = "bar"
    segment_chart.title = "Revenue by Customer Segment"

    segment_values = Reference(
        data_ws,
        min_col=6,
        min_row=1,
        max_row=len(segment_data) + 1
    )

    segment_categories = Reference(
        data_ws,
        min_col=5,
        min_row=2,
        max_row=len(segment_data) + 1
    )

    segment_chart.add_data(
        segment_values,
        titles_from_data=True
    )

    segment_chart.set_categories(
        segment_categories
    )

    segment_chart.height = 7
    segment_chart.width = 12

    dashboard.add_chart(
        segment_chart,
        "A30"
    )

    # ========================================================
    # CHART 4 — CATEGORY REVENUE
    # ========================================================

    category_chart = BarChart()

    category_chart.title = (
        "Revenue by Product Category"
    )

    category_values = Reference(
        data_ws,
        min_col=9,
        min_row=1,
        max_row=len(category_data) + 1
    )

    category_categories = Reference(
        data_ws,
        min_col=8,
        min_row=2,
        max_row=len(category_data) + 1
    )

    category_chart.add_data(
        category_values,
        titles_from_data=True
    )

    category_chart.set_categories(
        category_categories
    )

    category_chart.height = 7
    category_chart.width = 12

    dashboard.add_chart(
        category_chart,
        "G30"
    )

    # ========================================================
    # CHART 5 — DELIVERY PERFORMANCE
    # ========================================================

    delivery_chart = BarChart()

    delivery_chart.title = "Shipment Performance"

    delivery_values = Reference(
        data_ws,
        min_col=12,
        min_row=1,
        max_row=len(delivery_data) + 1
    )

    delivery_categories = Reference(
        data_ws,
        min_col=11,
        min_row=2,
        max_row=len(delivery_data) + 1
    )

    delivery_chart.add_data(
        delivery_values,
        titles_from_data=True
    )

    delivery_chart.set_categories(
        delivery_categories
    )

    delivery_chart.height = 7
    delivery_chart.width = 12

    dashboard.add_chart(
        delivery_chart,
        "A46"
    )

    # ========================================================
    # CHART 6 — WAREHOUSE SHIPMENT VOLUME
    # ========================================================

    warehouse_chart = BarChart()

    warehouse_chart.type = "bar"
    warehouse_chart.title = (
        "Warehouse Shipment Volume"
    )

    top_n = min(
        len(warehouse_data),
        10
    )

    warehouse_values = Reference(
        data_ws,
        min_col=15,
        min_row=1,
        max_row=top_n + 1
    )

    warehouse_categories = Reference(
        data_ws,
        min_col=14,
        min_row=2,
        max_row=top_n + 1
    )

    warehouse_chart.add_data(
        warehouse_values,
        titles_from_data=True
    )

    warehouse_chart.set_categories(
        warehouse_categories
    )

    warehouse_chart.height = 7
    warehouse_chart.width = 12

    dashboard.add_chart(
        warehouse_chart,
        "G46"
    )

    # ========================================================
    # CHART DATA FORMATTING
    # ========================================================

    for cell in data_ws[1]:

        cell.font = Font(
            bold=True
        )

    # Hide technical data sheet
    data_ws.sheet_state = "hidden"

    # ========================================================
    # PAGE SETTINGS
    # ========================================================

    dashboard.sheet_properties.pageSetUpPr.fitToPage = True

    dashboard.page_setup.orientation = "landscape"

    dashboard.page_setup.fitToWidth = 1

    dashboard.page_setup.fitToHeight = 0

    # ========================================================
    # SAVE
    # ========================================================

    print("\nSaving dashboard...")

    wb.save(
        OUTPUT_FILE
    )

    print("\n" + "=" * 70)
    print("LOGIX EXECUTIVE DASHBOARD CREATED SUCCESSFULLY")
    print("=" * 70)

    print("\nFile:")

    print(
        os.path.abspath(
            OUTPUT_FILE
        )
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()