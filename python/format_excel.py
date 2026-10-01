import os
import shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


# ============================================================
# LOGIX — Professional Excel Formatter
# ============================================================

INPUT_FILE = os.path.join(
    "excel",
    "LOGIX_Excel_Analysis.xlsx"
)

BACKUP_FILE = os.path.join(
    "excel",
    "LOGIX_Excel_Analysis_backup.xlsx"
)

OUTPUT_FILE = os.path.join(
    "excel",
    "LOGIX_Excel_Analysis_Professional.xlsx"
)


# ============================================================
# Sheet Configuration
# ============================================================

SHEET_CONFIG = {
    "Executive_KPI": {
        "currency": [
            "total_revenue",
            "total_logistics_cost",
            "total_refunds"
        ],
        "percentage": [
            "on_time_rate"
        ]
    },

    "Sales_Analysis": {
        "currency": [
            "unit_price",
            "discount_amount",
            "line_total"
        ],
        "percentage": [
            "discount_percent"
        ]
    },

    "Customer_Analysis": {
        "currency": [
            "total_revenue",
            "aov"
        ]
    },

    "Delivery_Analysis": {
        "currency": [
            "shipping_cost"
        ],
        "percentage": [],
        "decimal": [
            "distance_km",
            "delivery_days"
        ]
    },

    "Warehouse_Analysis": {
        "currency": [
            "total_logistics_cost"
        ],
        "percentage": [
            "on_time_rate"
        ],
        "decimal": [
            "avg_distance_km"
        ]
    },

    "Inventory_Analysis": {
        "currency": [
            "unit_cost",
            "inventory_value"
        ]
    },

    "Returns_Analysis": {
        "currency": [
            "return_amount",
            "refund_amount"
        ]
    },

    "Payment_Analysis": {
        "currency": [
            "amount"
        ]
    },

    "Monthly_Sales": {
        "currency": [
            "revenue",
            "aov"
        ]
    },

    "Fleet_Analysis": {
        "decimal": [
            "capacity_kg",
            "rating"
        ]
    }
}


# ============================================================
# Styling
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
    size=11
)

TITLE_FONT = Font(
    bold=True,
    size=14
)

THIN_BORDER = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2"
    )
)

CENTER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center"
)

LEFT_ALIGNMENT = Alignment(
    horizontal="left",
    vertical="center"
)


# ============================================================
# Helper Functions
# ============================================================

def normalize_header(value):
    return str(value).strip().lower()


def get_headers(ws):
    return {
        normalize_header(cell.value): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def format_currency(cell):
    cell.number_format = '₹#,##0.00'


def format_percentage(cell):
    cell.number_format = '0.00%'


def format_decimal(cell):
    cell.number_format = '#,##0.00'


def apply_conditional_formatting(ws, headers):

    # --------------------------------------------------------
    # Status columns
    # --------------------------------------------------------

    for header, col in headers.items():

        if "status" in header or "performance" in header:

            letter = get_column_letter(col)

            data_range = f"{letter}2:{letter}{ws.max_row}"

            # Highlight negative operational states
            if any(
                word in header
                for word in [
                    "status",
                    "performance"
                ]
            ):

                ws.conditional_formatting.add(
                    data_range,
                    FormulaRule(
                        formula=[
                            f'OR(ISNUMBER(SEARCH("Failed",{letter}2)),'
                            f'ISNUMBER(SEARCH("Delayed",{letter}2)),'
                            f'ISNUMBER(SEARCH("Rejected",{letter}2)))'
                        ],
                        fill=PatternFill(
                            "solid",
                            fgColor="F4CCCC"
                        )
                    )
                )

    # --------------------------------------------------------
    # Stock status
    # --------------------------------------------------------

    if "stock_status" in headers:

        col = get_column_letter(
            headers["stock_status"]
        )

        data_range = f"{col}2:{col}{ws.max_row}"

        ws.conditional_formatting.add(
            data_range,
            FormulaRule(
                formula=[
                    f'OR(ISNUMBER(SEARCH("Low",{col}2)),'
                    f'ISNUMBER(SEARCH("Critical",{col}2)))'
                ],
                fill=PatternFill(
                    "solid",
                    fgColor="F4CCCC"
                )
            )
        )


# ============================================================
# Main Formatting
# ============================================================

def main():

    print("=" * 70)
    print("LOGIX — PROFESSIONAL EXCEL FORMATTER")
    print("=" * 70)

    if not os.path.exists(INPUT_FILE):
        raise FileNotFoundError(
            f"Input Excel file not found: {INPUT_FILE}"
        )

    # --------------------------------------------------------
    # Backup original workbook
    # --------------------------------------------------------

    print("\nCreating backup...")

    shutil.copy2(
        INPUT_FILE,
        BACKUP_FILE
    )

    print(
        f"Backup created: {BACKUP_FILE}"
    )

    # --------------------------------------------------------
    # Load workbook
    # --------------------------------------------------------

    wb = load_workbook(INPUT_FILE)

    print("\nFormatting worksheets...")

    for ws in wb.worksheets:

        print(f"\nProcessing: {ws.title}")

        # ----------------------------------------------------
        # Freeze top row
        # ----------------------------------------------------

        ws.freeze_panes = "A2"

        # ----------------------------------------------------
        # Header formatting
        # ----------------------------------------------------

        for cell in ws[1]:

            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER_ALIGNMENT
            cell.border = THIN_BORDER

        ws.row_dimensions[1].height = 24

        # ----------------------------------------------------
        # Headers
        # ----------------------------------------------------

        headers = get_headers(ws)

        config = SHEET_CONFIG.get(
            ws.title,
            {}
        )

        currency_columns = config.get(
            "currency",
            []
        )

        percentage_columns = config.get(
            "percentage",
            []
        )

        decimal_columns = config.get(
            "decimal",
            []
        )

        # ----------------------------------------------------
        # Data formatting
        # ----------------------------------------------------

        for header, col in headers.items():

            column_letter = get_column_letter(col)

            # Width
            max_length = 0

            for cell in ws[column_letter]:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            width = min(
                max(max_length + 2, 12),
                35
            )

            ws.column_dimensions[
                column_letter
            ].width = width

            # Currency
            if header in currency_columns:

                for row in range(
                    2,
                    ws.max_row + 1
                ):

                    format_currency(
                        ws.cell(row, col)
                    )

            # Percentage
            if header in percentage_columns:

                for row in range(
                    2,
                    ws.max_row + 1
                ):

                    format_percentage(
                        ws.cell(row, col)
                    )

            # Decimal
            if header in decimal_columns:

                for row in range(
                    2,
                    ws.max_row + 1
                ):

                    format_decimal(
                        ws.cell(row, col)
                    )

        # ----------------------------------------------------
        # Alignment
        # ----------------------------------------------------

        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="center"
                )

        # ----------------------------------------------------
        # Excel Table
        # ----------------------------------------------------

        if ws.max_row >= 2:

            table_name = (
                "tbl_"
                + ws.title.replace(
                    "_",
                    ""
                )
            )

            table_name = table_name[:240]

            # Remove existing tables if any
            ws.tables.clear()

            table = Table(
                displayName=table_name,
                ref=ws.dimensions
            )

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )

            table.tableStyleInfo = style

            ws.add_table(table)

        # ----------------------------------------------------
        # Conditional Formatting
        # ----------------------------------------------------

        apply_conditional_formatting(
            ws,
            headers
        )

        # ----------------------------------------------------
        # AutoFilter
        # ----------------------------------------------------

        ws.auto_filter.ref = ws.dimensions

        # ----------------------------------------------------
        # Print settings
        # ----------------------------------------------------

        ws.sheet_view.showGridLines = False

        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ========================================================
    # Save
    # ========================================================

    wb.save(OUTPUT_FILE)

    print("\n" + "=" * 70)
    print("FORMATTING COMPLETED SUCCESSFULLY")
    print("=" * 70)

    print(
        f"\nProfessional workbook:"
    )

    print(
        os.path.abspath(OUTPUT_FILE)
    )

    print(
        f"\nBackup workbook:"
    )

    print(
        os.path.abspath(BACKUP_FILE)
    )


if __name__ == "__main__":
    main()