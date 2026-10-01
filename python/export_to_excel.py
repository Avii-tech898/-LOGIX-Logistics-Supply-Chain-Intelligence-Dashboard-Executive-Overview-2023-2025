import os
import pandas as pd
import mysql.connector

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# LOGIX — Optimized MySQL → Professional Excel Export
# ============================================================

DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "avanish@123",
    "database": "logix"
}

OUTPUT_DIR = "excel"

OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    "LOGIX_Excel_Analysis_Professional.xlsx"
)


VIEWS = {
    "Executive_KPI": "vw_executive_kpi",
    "Sales_Analysis": "vw_sales_analysis",
    "Customer_Analysis": "vw_customer_analysis",
    "Delivery_Analysis": "vw_delivery_performance",
    "Warehouse_Analysis": "vw_warehouse_performance",
    "Inventory_Analysis": "vw_inventory_analysis",
    "Returns_Analysis": "vw_returns_analysis",
    "Payment_Analysis": "vw_payment_analysis",
    "Monthly_Sales": "vw_monthly_sales",
    "Fleet_Analysis": "vw_fleet_analysis"
}


# ============================================================
# Formatting configuration
# ============================================================

CURRENCY_COLUMNS = {
    "Executive_KPI": [
        "total_revenue",
        "total_logistics_cost",
        "total_refunds"
    ],

    "Sales_Analysis": [
        "unit_price",
        "discount_amount",
        "line_total"
    ],

    "Customer_Analysis": [
        "total_revenue",
        "aov"
    ],

    "Delivery_Analysis": [
        "shipping_cost"
    ],

    "Warehouse_Analysis": [
        "total_logistics_cost"
    ],

    "Inventory_Analysis": [
        "unit_cost",
        "inventory_value"
    ],

    "Returns_Analysis": [
        "return_amount",
        "refund_amount"
    ],

    "Payment_Analysis": [
        "amount"
    ],

    "Monthly_Sales": [
        "revenue",
        "aov"
    ]
}


PERCENTAGE_COLUMNS = {
    "Executive_KPI": [
        "on_time_rate"
    ],

    "Sales_Analysis": [
        "discount_percent"
    ],

    "Warehouse_Analysis": [
        "on_time_rate"
    ]
}


DATE_KEYWORDS = [
    "date",
    "datetime"
]


# ============================================================
# Excel styles
# ============================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
    size=11
)

HEADER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center"
)


# ============================================================
# Helper
# ============================================================

def clean_name(name):
    return str(name).strip().lower()


def apply_number_formats(ws, sheet_name):

    headers = {}

    for cell in ws[1]:
        if cell.value is not None:
            headers[
                clean_name(cell.value)
            ] = cell.column

    currency_columns = [
        clean_name(x)
        for x in CURRENCY_COLUMNS.get(
            sheet_name,
            []
        )
    ]

    percentage_columns = [
        clean_name(x)
        for x in PERCENTAGE_COLUMNS.get(
            sheet_name,
            []
        )
    ]

    # --------------------------------------------------------
    # Only apply formats to configured columns.
    # No full workbook scan.
    # --------------------------------------------------------

    for column_name in currency_columns:

        if column_name not in headers:
            continue

        col = headers[column_name]

        for row in range(
            2,
            ws.max_row + 1
        ):
            ws.cell(
                row=row,
                column=col
            ).number_format = '₹#,##0.00'

    for column_name in percentage_columns:

        if column_name not in headers:
            continue

        col = headers[column_name]

        for row in range(
            2,
            ws.max_row + 1
        ):
            ws.cell(
                row=row,
                column=col
            ).number_format = '0.00%'


def apply_date_formats(ws):

    for cell in ws[1]:

        if cell.value is None:
            continue

        header = clean_name(cell.value)

        if any(
            keyword in header
            for keyword in DATE_KEYWORDS
        ):

            col = cell.column

            for row in range(
                2,
                ws.max_row + 1
            ):
                ws.cell(
                    row=row,
                    column=col
                ).number_format = "yyyy-mm-dd"


def set_column_widths(ws):

    # Efficient fixed widths based on column type.
    # Avoid scanning hundreds of thousands of rows.

    for cell in ws[1]:

        if cell.value is None:
            continue

        header = clean_name(cell.value)
        col_letter = get_column_letter(
            cell.column
        )

        if "id" in header:
            width = 18

        elif "name" in header:
            width = 25

        elif "description" in header:
            width = 35

        elif "address" in header:
            width = 30

        elif "email" in header:
            width = 30

        elif "phone" in header:
            width = 18

        elif "status" in header:
            width = 20

        elif "date" in header:
            width = 16

        else:
            width = 18

        ws.column_dimensions[
            col_letter
        ].width = width


def format_header(ws):

    for cell in ws[1]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    ws.row_dimensions[1].height = 24


def create_table(ws, sheet_name):

    if ws.max_row < 2:
        return

    table_name = (
        "tbl_"
        + sheet_name.replace("_", "")
    )

    table = Table(
        displayName=table_name,
        ref=ws.dimensions
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    table.tableStyleInfo = style

    ws.add_table(table)


# ============================================================
# Main
# ============================================================

def main():

    print("=" * 70)
    print("LOGIX — OPTIMIZED PROFESSIONAL EXCEL EXPORT")
    print("=" * 70)

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    print("\nConnecting to MySQL...")

    connection = mysql.connector.connect(
        host=DB_CONFIG["host"],
        user=DB_CONFIG["user"],
        password=DB_CONFIG["password"],
        database=DB_CONFIG["database"]
    )

    print("MySQL connection successful.")

    # --------------------------------------------------------
    # Export data
    # --------------------------------------------------------

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        for sheet_name, view_name in VIEWS.items():

            print(
                f"\nExporting: {view_name}"
            )

            query = f"SELECT * FROM {view_name}"

            df = pd.read_sql(
                query,
                connection
            )

            print(
                f"Rows: {len(df):,}"
            )

            print(
                f"Columns: {len(df.columns)}"
            )

            df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

            print(
                f"Written → {sheet_name}"
            )

    connection.close()

    print("\nMySQL connection closed.")

    # --------------------------------------------------------
    # Apply lightweight formatting
    # --------------------------------------------------------

    print("\nApplying Excel formatting...")

    wb = load_workbook(
        OUTPUT_FILE
    )

    for ws in wb.worksheets:

        print(
            f"Formatting → {ws.title}"
        )

        # Freeze header
        ws.freeze_panes = "A2"

        # Hide gridlines
        ws.sheet_view.showGridLines = False

        # Header
        format_header(ws)

        # Widths
        set_column_widths(ws)

        # Number formats
        apply_number_formats(
            ws,
            ws.title
        )

        # Dates
        apply_date_formats(ws)

        # Table
        create_table(
            ws,
            ws.title
        )

        # Print settings
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    wb.save(
        OUTPUT_FILE
    )

    print("\n" + "=" * 70)
    print("LOGIX EXCEL EXPORT COMPLETED")
    print("=" * 70)

    print(
        "\nProfessional Excel file:"
    )

    print(
        os.path.abspath(
            OUTPUT_FILE
        )
    )


if __name__ == "__main__":
    main()