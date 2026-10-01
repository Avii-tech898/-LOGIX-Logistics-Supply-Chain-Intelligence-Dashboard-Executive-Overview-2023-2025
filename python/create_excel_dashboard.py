import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# LOGIX — Executive Dashboard
# ============================================================

INPUT_FILE = os.path.join(
    "excel",
    "LOGIX_Excel_Analysis_Professional.xlsx"
)

OUTPUT_FILE = os.path.join(
    "excel",
    "LOGIX_Executive_Dashboard.xlsx"
)


# ============================================================
# Theme
# ============================================================

DARK_BLUE = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"
GREEN = "E2F0D9"
RED = "F4CCCC"
GOLD = "FFF2CC"


# ============================================================
# Styles
# ============================================================

title_fill = PatternFill(
    "solid",
    fgColor=DARK_BLUE
)

kpi_fill = PatternFill(
    "solid",
    fgColor=LIGHT_BLUE
)

section_fill = PatternFill(
    "solid",
    fgColor=BLUE
)

white_font = Font(
    color=WHITE,
    bold=True
)

title_font = Font(
    color=WHITE,
    bold=True,
    size=20
)

kpi_label_font = Font(
    bold=True,
    size=10
)

kpi_value_font = Font(
    bold=True,
    size=16
)

thin_border = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2")
)


# ============================================================
# Main
# ============================================================

def main():

    print("=" * 70)
    print("LOGIX — EXECUTIVE DASHBOARD")
    print("=" * 70)

    if not os.path.exists(INPUT_FILE):
        raise FileNotFoundError(
            f"Workbook not found: {INPUT_FILE}"
        )

    print("\nLoading professional workbook...")

    wb = load_workbook(
        INPUT_FILE
    )

    # --------------------------------------------------------
    # Remove old dashboard if exists
    # --------------------------------------------------------

    if "Executive_Dashboard" in wb.sheetnames:

        del wb["Executive_Dashboard"]

    # --------------------------------------------------------
    # Create dashboard
    # --------------------------------------------------------

    ws = wb.create_sheet(
        "Executive_Dashboard",
        0
    )

    # --------------------------------------------------------
    # General setup
    # --------------------------------------------------------

    ws.sheet_view.showGridLines = False

    ws.freeze_panes = "A5"

    # Column widths
    for col in range(1, 13):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = 15

    # --------------------------------------------------------
    # Dashboard Title
    # --------------------------------------------------------

    ws.merge_cells(
        "A1:L2"
    )

    title = ws["A1"]

    title.value = (
        "LOGIX — LOGISTICS & SUPPLY CHAIN "
        "INTELLIGENCE DASHBOARD"
    )

    title.fill = title_fill
    title.font = title_font
    title.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    # --------------------------------------------------------
    # Subtitle
    # --------------------------------------------------------

    ws.merge_cells(
        "A3:L3"
    )

    subtitle = ws["A3"]

    subtitle.value = (
        "Executive Business Performance Overview"
    )

    subtitle.font = Font(
        italic=True,
        size=10
    )

    subtitle.alignment = Alignment(
        horizontal="center"
    )

    # --------------------------------------------------------
    # KPI definitions
    # --------------------------------------------------------

    kpis = [
        ("Total Orders", "total_orders"),
        ("Total Customers", "total_customers"),
        ("Total Revenue", "total_revenue"),
        ("Logistics Cost", "total_logistics_cost"),
        ("Failed Shipments", "failed_shipments"),
        ("Total Returns", "total_returns"),
        ("Total Refunds", "total_refunds"),
        ("On-Time Delivery", "on_time_rate")
    ]

    # --------------------------------------------------------
    # Read Executive KPI source
    # --------------------------------------------------------

    source_ws = wb["Executive_KPI"]

    headers = {
        str(cell.value).strip(): cell.column
        for cell in source_ws[1]
        if cell.value is not None
    }

    values = {}

    for name, col in headers.items():

        values[name] = source_ws.cell(
            row=2,
            column=col
        ).value

    # --------------------------------------------------------
    # KPI Cards
    # --------------------------------------------------------

    positions = [
        ("A5:C7"),
        ("D5:F7"),
        ("G5:I7"),
        ("J5:L7"),

        ("A9:C11"),
        ("D9:F11"),
        ("G9:I11"),
        ("J9:L11")
    ]

    for index, (
        label,
        field
    ) in enumerate(kpis):

        cell_range = positions[index]

        ws.merge_cells(
            cell_range
        )

        start_cell = ws[
            cell_range.split(":")[0]
        ]

        start_cell.value = (
            f"{label}\n"
        )

        start_cell.fill = kpi_fill

        start_cell.font = kpi_label_font

        start_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        start_cell.border = thin_border

        # Value row below the label area
        start_row = start_cell.row
        start_col = start_cell.column

        value_row = start_row + 1

        # We use the top-left cell for display
        start_cell.value = (
            f"{label}\n"
            f"{values.get(field, 0)}"
        )

        start_cell.font = Font(
            bold=True,
            size=13
        )

    # --------------------------------------------------------
    # Format KPI values
    # --------------------------------------------------------

    currency_fields = {
        "total_revenue",
        "total_logistics_cost",
        "total_refunds"
    }

    percentage_fields = {
        "on_time_rate"
    }

    for index, (
        label,
        field
    ) in enumerate(kpis):

        cell_range = positions[index]

        cell = ws[
            cell_range.split(":")[0]
        ]

        value = values.get(
            field,
            0
        )

        if field in currency_fields:

            cell.value = (
                f"{label}\n"
                f"₹{value:,.2f}"
            )

        elif field in percentage_fields:

            cell.value = (
                f"{label}\n"
                f"{value:.2f}%"
    )

    else:

        cell.value = (
                f"{label}\n"
                f"{value:,}"
            )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # --------------------------------------------------------
    # Section headers
    # --------------------------------------------------------

    sections = [
        ("A13:F13", "SALES PERFORMANCE"),
        ("G13:L13", "DELIVERY PERFORMANCE"),
        ("A25:F25", "CUSTOMER & REVENUE"),
        ("G25:L25", "OPERATIONS")
    ]

    for cell_range, text in sections:

        ws.merge_cells(
            cell_range
        )

        cell = ws[
            cell_range.split(":")[0]
        ]

        cell.value = text
        cell.fill = section_fill
        cell.font = white_font

        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    print("\nSaving dashboard...")

    wb.save(
        OUTPUT_FILE
    )

    print("\n" + "=" * 70)
    print("EXECUTIVE DASHBOARD FOUNDATION CREATED")
    print("=" * 70)

    print(
        f"\nFile:"
    )

    print(
        os.path.abspath(
            OUTPUT_FILE
        )
    )


if __name__ == "__main__":
    main()